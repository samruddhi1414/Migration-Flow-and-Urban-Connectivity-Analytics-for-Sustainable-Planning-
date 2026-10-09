import pandas as pd
import openpyxl
import os

raw_dir = r"C:\Users\Amey\Desktop\raw"

# Let's inspect village columns in detail
wb = openpyxl.load_workbook(os.path.join(raw_dir, "village_dchb.xlsx"), read_only=True)
ws = wb['Village_Data_2700']

# Get headers
headers = [cell for cell in next(ws.iter_rows(values_only=True))]
print(f"Total village columns: {len(headers)}")

# Check district distribution and sample row
dist_counts = {}
sample_rows = []
for i, row in enumerate(ws.iter_rows(values_only=True)):
    if i == 0:
        continue
    dcode = str(row[2]).zfill(3) if row[2] is not None else "UNKNOWN"
    dist_counts[dcode] = dist_counts.get(dcode, 0) + 1
    if i <= 3:
        sample_rows.append(row)

print(f"Total village rows: {sum(dist_counts.values())}")
print("Districts in village data:")
for dcode, count in sorted(dist_counts.items()):
    print(f"  District {dcode}: {count} villages")

wb.close()
